from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


START_DATE = "2026-09-01"
END_DATE = "2026-09-30"


def load_env(path: Path) -> None:
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip())


def as_number(value: Any) -> float:
    try:
        number = float(value or 0)
        return round(number, 6)
    except (TypeError, ValueError):
        return 0.0


def as_datetime(value: Any, epoch: datetime) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted if isinstance(converted, datetime) else datetime.combine(converted, datetime.min.time())
        except (TypeError, ValueError, OverflowError):
            return None
    return None


def iso_date(value: Any, epoch: datetime) -> str:
    parsed = as_datetime(value, epoch)
    return parsed.date().isoformat() if parsed else ""


def iso_timestamp(value: Any, epoch: datetime) -> str | None:
    parsed = as_datetime(value, epoch)
    if not parsed:
        return None
    return parsed.replace(tzinfo=timezone.utc).isoformat().replace("+00:00", "Z")


def in_scope(value: Any, epoch: datetime) -> bool:
    parsed = iso_date(value, epoch)
    return START_DATE <= parsed <= END_DATE


def text(value: Any) -> str:
    return str(value or "").strip()


def fee_exempt_product(value: Any) -> bool:
    normalized = " ".join("".join(character.lower() if character.isalnum() else " " for character in text(value)).split())
    return "slimboost" in normalized.replace(" ", "") or "detox" in normalized


def stable_integer(*parts: str) -> str:
    digest = hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()
    return str(8_000_000_000_000_000 + int(digest[:12], 16) % 999_999_999_999_999)


class Supabase:
    def __init__(self, url: str, key: str) -> None:
        normalized = url.rstrip("/")
        self.base = normalized if normalized.endswith("/rest/v1") else f"{normalized}/rest/v1"
        self.headers = {"apikey": key, "Authorization": f"Bearer {key}"}

    def request(self, method: str, path: str, body: Any = None, prefer: str = "") -> Any:
        payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
        headers = dict(self.headers)
        if payload is not None:
            headers["Content-Type"] = "application/json"
        if prefer:
            headers["Prefer"] = prefer
        request = urllib.request.Request(f"{self.base}/{path}", data=payload, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                content = response.read()
                return json.loads(content) if content else None
        except urllib.error.HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"Supabase {method} {path} failed ({error.code}): {detail}") from error

    def select(self, table: str, filters: str = "") -> list[dict[str, Any]]:
        path = f"{table}?select=*{('&' + filters) if filters else ''}&limit=10000"
        result = self.request("GET", path)
        return result if isinstance(result, list) else []

    def delete(self, table: str, filters: str) -> None:
        self.request("DELETE", f"{table}?{filters}", prefer="return=minimal")

    def insert(self, table: str, rows: list[dict[str, Any]], conflict: str = "") -> None:
        if not rows:
            return
        suffix = f"?on_conflict={conflict}" if conflict else ""
        prefer = "resolution=merge-duplicates,return=minimal" if conflict else "return=minimal"
        for offset in range(0, len(rows), 250):
            self.request("POST", f"{table}{suffix}", rows[offset:offset + 250], prefer=prefer)


def parse_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    required = {"ADS", "Orders", "Total sales by order"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(sorted(missing))}")
    epoch = workbook.epoch
    now = datetime.now(timezone.utc).isoformat()

    ads_rows: list[dict[str, Any]] = []
    cogs_by_date: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ads_sheet = workbook["ADS"]
    for row_number, row in enumerate(ads_sheet.iter_rows(values_only=True), 1):
        values = list(row)
        summary_date = values[1] if len(values) > 1 else None
        if in_scope(summary_date, epoch):
            report_date = iso_date(summary_date, epoch)
            ads_rows.append({
                "report_date": report_date,
                "timezone": "America/New_York",
                "currency": "USD",
                "meta_cost": as_number(values[2]),
                "google_cost": as_number(values[3]),
                "tiktok_cost": as_number(values[4]),
                "cogs": as_number(values[5]),
                "shipping": as_number(values[6]),
                "fulfillment": as_number(values[7]),
                "processing": as_number(values[8]),
                "fetched_at": now,
            })

        detail_date = values[11] if len(values) > 11 else None
        if in_scope(detail_date, epoch):
            report_date = iso_date(detail_date, epoch)
            subtotal = as_number(values[16])
            shipping = as_number(values[17])
            fulfillment = as_number(values[18])
            processing_supliful = as_number(values[19])
            processing_shopify = as_number(values[20])
            cogs_by_date[report_date].append({
                "id": f"excel-ads-{row_number}",
                "date": report_date,
                "order": text(values[12]),
                "product": text(values[13]),
                "qty": as_number(values[14]),
                "unit_price": as_number(values[15]),
                "subtotal": subtotal,
                "shipping": shipping,
                "fulfillment_supliful": fulfillment,
                "processing_supliful": processing_supliful,
                "processing_shopify": processing_shopify,
                "total": as_number(values[21]) or round(subtotal + shipping + fulfillment + processing_supliful + processing_shopify, 6),
            })

    orders_by_date: dict[str, list[dict[str, Any]]] = defaultdict(list)
    order_rows_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    tiktok_orders: set[str] = set()
    orders_sheet = workbook["Orders"]
    for row_number, row in enumerate(orders_sheet.iter_rows(min_row=2, values_only=True), 2):
        values = list(row)
        created_at = values[15] if len(values) > 15 else None
        order_name = text(values[0])
        if "tiktokw.us" in text(values[1]).lower():
            tiktok_orders.add(order_name)
        if not in_scope(created_at, epoch):
            continue
        report_date = iso_date(created_at, epoch)
        order_row = {
            "shopify_order_id": stable_integer("order", order_name),
            "shopify_lineitem_id": stable_integer("line", order_name, str(row_number)),
            "order_date": report_date,
            "name": order_name,
            "email": text(values[1]) or None,
            "financial_status": text(values[2]),
            "paid_at": iso_timestamp(values[3], epoch),
            "lineitem_quantity": int(as_number(values[16])),
            "lineitem_name": text(values[17]),
            "lineitem_price": as_number(values[18]),
            "lineitem_compare_at_price": None if values[19] in (None, "") else as_number(values[19]),
            "lineitem_sku": text(values[20]) or None,
            "fetched_at": now,
        }
        orders_by_date[report_date].append(order_row)
        order_rows_by_name[order_name].append(order_row)

    # The workbook does not include a dedicated sales-channel column, but
    # TikTok Shop orders use Shopify's @scs.tiktokw.us relay email domain.
    # Product exemptions apply independently of the order channel.
    adjusted_costs_by_date: dict[str, dict[str, float]] = defaultdict(lambda: {"cogs": 0.0, "shipping": 0.0, "fulfillment": 0.0, "processing": 0.0})
    for report_date, rows in cogs_by_date.items():
        for row in rows:
            tiktok_order = row["order"] in tiktok_orders
            exempt_product = fee_exempt_product(row["product"])
            if tiktok_order:
                row["unit_price"] = 0.01
                row["subtotal"] = round(as_number(row["qty"]) * 0.01, 6)
            if exempt_product or tiktok_order:
                row["shipping"] = 0.0
                row["fulfillment_supliful"] = 0.0
                row["processing_supliful"] = 0.0
                row["processing_shopify"] = 0.0
                row["total"] = round(as_number(row["subtotal"]), 6)
            adjusted_costs_by_date[report_date]["cogs"] += as_number(row["subtotal"])
            adjusted_costs_by_date[report_date]["shipping"] += as_number(row["shipping"])
            adjusted_costs_by_date[report_date]["fulfillment"] += as_number(row["fulfillment_supliful"])
            adjusted_costs_by_date[report_date]["processing"] += as_number(row["processing_supliful"])
    for row in ads_rows:
        adjusted = adjusted_costs_by_date[row["report_date"]]
        row["cogs"] = round(adjusted["cogs"], 6)
        row["shipping"] = round(adjusted["shipping"], 6)
        row["fulfillment"] = round(adjusted["fulfillment"], 6)
        row["processing"] = round(adjusted["processing"], 6)

    raw_sales_by_date: dict[str, list[dict[str, Any]]] = defaultdict(list)
    sales_sheet = workbook["Total sales by order"]
    for row_number, row in enumerate(sales_sheet.iter_rows(min_row=2, values_only=True), 2):
        values = list(row)
        if not in_scope(values[0], epoch):
            continue
        report_date = iso_date(values[0], epoch)
        raw_sales_by_date[report_date].append({
            "row_number": row_number,
            "day": report_date,
            "sale_id": text(values[1]),
            "order_name": text(values[2]),
            "product_title": text(values[3]),
            "line_gross_sales": as_number(values[4]),
            "line_discounts": as_number(values[5]),
            "line_returns": as_number(values[6]),
            "line_net_sales": as_number(values[7]),
            "line_shipping_charges": as_number(values[8]),
            "line_return_fees": as_number(values[9]),
            "line_taxes": as_number(values[10]),
            "line_total_sales": as_number(values[11]),
        })

    sales_by_date: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for report_date, source_rows in raw_sales_by_date.items():
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for source_row in source_rows:
            grouped[source_row["order_name"]].append(source_row)
        for order_name, group in grouped.items():
            order_lines = order_rows_by_name.get(order_name, [])
            product_line_index = 0
            totals = {
                "gross_sales": round(sum(row["line_gross_sales"] for row in group), 6),
                "discounts": round(sum(row["line_discounts"] for row in group), 6),
                "returns": round(sum(row["line_returns"] for row in group), 6),
                "net_sales": round(sum(row["line_net_sales"] for row in group), 6),
                "shipping_charges": round(sum(row["line_shipping_charges"] for row in group), 6),
                "total_sales": round(sum(row["line_total_sales"] for row in group), 6),
            }
            for index, source_row in enumerate(group):
                matched_order_line = None
                if source_row["product_title"]:
                    if product_line_index < len(order_lines):
                        matched_order_line = order_lines[product_line_index]
                    product_line_index += 1
                report_row = {
                    "id": f"excel-{source_row['sale_id'] or source_row['row_number']}",
                    **{key: value for key, value in source_row.items() if key != "row_number"},
                    "name": order_name if index == 0 else "",
                    "gross_sales": totals["gross_sales"] if index == 0 else None,
                    "discounts": totals["discounts"] if index == 0 else None,
                    "returns": totals["returns"] if index == 0 else None,
                    "net_sales": totals["net_sales"] if index == 0 else None,
                    "shipping_charges": totals["shipping_charges"] if index == 0 else None,
                    "total_sales": totals["total_sales"] if index == 0 else None,
                    "qty": matched_order_line["lineitem_quantity"] if matched_order_line else 0,
                    "sales": source_row["line_net_sales"],
                    "product_name": matched_order_line["lineitem_name"] if matched_order_line else source_row["product_title"],
                }
                sales_by_date[report_date].append(report_row)

    ads_dates = [row["report_date"] for row in ads_rows]
    expected_days = (date.fromisoformat(END_DATE) - date.fromisoformat(START_DATE)).days + 1
    if len(ads_rows) != expected_days or len(set(ads_dates)) != expected_days:
        raise ValueError(f"Expected exactly {expected_days} unique ADS rows; found {len(ads_rows)} rows and {len(set(ads_dates))} dates")

    return {
        "ads": ads_rows,
        "cogs": [{"report_date": report_date, "report_data": rows, "fetched_at": now} for report_date, rows in sorted(cogs_by_date.items())],
        "orders": [row for report_date in sorted(orders_by_date) for row in orders_by_date[report_date]],
        "sales": [{"report_date": report_date, "report_data": {"date": report_date, "fetchedAt": now, "rows": rows}, "fetched_at": now} for report_date, rows in sorted(sales_by_date.items())],
    }


def scoped_filters(date_field: str) -> str:
    return f"{date_field}=gte.{START_DATE}&{date_field}=lte.{END_DATE}"


def backup_current(database: Supabase) -> dict[str, list[dict[str, Any]]]:
    return {
        "supplements_ads_reports": database.select("supplements_ads_reports", scoped_filters("report_date")),
        "shopify_cogs_reports": database.select("shopify_cogs_reports", scoped_filters("report_date")),
        "shopify_sales_reports": database.select("shopify_sales_reports", scoped_filters("report_date")),
        "shopify_order_line_items": database.select("shopify_order_line_items", scoped_filters("order_date")),
        "shopify_sales_history": database.select("shopify_sales_history", "report_key=eq.all-time"),
    }


def replace_september(database: Supabase, parsed: dict[str, Any]) -> None:
    database.delete("supplements_ads_reports", scoped_filters("report_date"))
    database.delete("shopify_cogs_reports", scoped_filters("report_date"))
    database.delete("shopify_sales_reports", scoped_filters("report_date"))
    database.delete("shopify_order_line_items", scoped_filters("order_date"))
    database.insert("supplements_ads_reports", parsed["ads"], "report_date")
    database.insert("shopify_cogs_reports", parsed["cogs"], "report_date")
    database.insert("shopify_sales_reports", parsed["sales"], "report_date")
    database.insert("shopify_order_line_items", parsed["orders"], "shopify_order_id,shopify_lineitem_id")


def restore_backup(database: Supabase, backup: dict[str, list[dict[str, Any]]]) -> None:
    database.delete("supplements_ads_reports", scoped_filters("report_date"))
    database.delete("shopify_cogs_reports", scoped_filters("report_date"))
    database.delete("shopify_sales_reports", scoped_filters("report_date"))
    database.delete("shopify_order_line_items", scoped_filters("order_date"))
    database.insert("supplements_ads_reports", backup["supplements_ads_reports"], "report_date")
    database.insert("shopify_cogs_reports", backup["shopify_cogs_reports"], "report_date")
    database.insert("shopify_sales_reports", backup["shopify_sales_reports"], "report_date")
    database.insert("shopify_order_line_items", backup["shopify_order_line_items"], "shopify_order_id,shopify_lineitem_id")
    database.delete("shopify_sales_history", "report_key=eq.all-time")
    database.insert("shopify_sales_history", backup["shopify_sales_history"], "report_key")


def rebuild_history(database: Supabase) -> None:
    snapshots = database.select("shopify_sales_reports", "order=report_date.asc")
    products: dict[str, dict[str, Any]] = {}
    for snapshot in snapshots:
        for row in snapshot.get("report_data", {}).get("rows", []):
            name = text(row.get("product_name") or row.get("product_title"))
            qty = as_number(row.get("qty"))
            sales = as_number(row.get("sales"))
            if not name or (qty == 0 and abs(sales) < 0.005):
                continue
            key = name.casefold()
            current = products.setdefault(key, {"product": name, "qty": 0, "sales_amount": 0.0})
            current["qty"] += qty
            current["sales_amount"] += sales
    historical_rows = sorted(({
        "product": row["product"],
        "qty": int(row["qty"]) if float(row["qty"]).is_integer() else row["qty"],
        "sales_amount": round(row["sales_amount"], 2),
    } for row in products.values()), key=lambda row: row["product"].casefold())
    fetched_at = datetime.now(timezone.utc).isoformat()
    database.insert("shopify_sales_history", [{
        "report_key": "all-time",
        "report_data": {"throughDate": END_DATE, "fetchedAt": fetched_at, "historicalRows": historical_rows},
        "fetched_at": fetched_at,
    }], "report_key")


def summary(parsed: dict[str, Any]) -> dict[str, Any]:
    cogs_rows = [row for snapshot in parsed["cogs"] for row in snapshot["report_data"]]
    sales_rows = [row for snapshot in parsed["sales"] for row in snapshot["report_data"]["rows"]]
    return {
        "ads_days": len(parsed["ads"]),
        "cogs_days": len(parsed["cogs"]),
        "cogs_rows": len(cogs_rows),
        "orders_rows": len(parsed["orders"]),
        "sales_days": len(parsed["sales"]),
        "sales_rows": len(sales_rows),
        "ads_cogs": round(sum(row["cogs"] for row in parsed["ads"]), 2),
        "detail_cogs": round(sum(as_number(row["subtotal"]) for row in cogs_rows), 2),
        "total_sales": round(sum(as_number(row["line_total_sales"]) for row in sales_rows), 2),
    }


def main() -> int:
    global END_DATE
    parser = argparse.ArgumentParser(description="Replace September supplement reporting data from the Dharma workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--through", default=END_DATE, help="Last September date to replace (YYYY-MM-DD).")
    parser.add_argument("--apply", action="store_true", help="Write the replacement data to Supabase.")
    parser.add_argument("--backup-dir", type=Path, default=Path("migration-backups"))
    args = parser.parse_args()
    if not (START_DATE <= args.through <= "2026-09-30"):
        raise ValueError("--through must be between 2026-09-01 and 2026-09-30")
    END_DATE = args.through

    load_env(Path(".env.local"))
    parsed = parse_workbook(args.workbook)
    print(json.dumps(summary(parsed), indent=2))
    if not args.apply:
        print("Dry run only. Use --apply to replace Supabase data.")
        return 0

    database = Supabase(os.environ["VITE_SUPABASE_URL"], os.environ["SUPABASE_SERVICE_ROLE_KEY"])
    backup = backup_current(database)
    args.backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = args.backup_dir / f"supplements-september-backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.json"
    backup_path.write_text(json.dumps(backup, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Backup written to {backup_path.resolve()}")
    try:
        replace_september(database, parsed)
        rebuild_history(database)
    except Exception:
        print("Migration failed; restoring the Supabase backup.", file=sys.stderr)
        restore_backup(database, backup)
        raise

    saved = backup_current(database)
    result = {
        "ads_days": len(saved["supplements_ads_reports"]),
        "cogs_days": len(saved["shopify_cogs_reports"]),
        "sales_days": len(saved["shopify_sales_reports"]),
        "orders_rows": len(saved["shopify_order_line_items"]),
        "history_snapshots": len(saved["shopify_sales_history"]),
    }
    print("Supabase verification:")
    print(json.dumps(result, indent=2))
    expected = {"ads_days": len(parsed["ads"]), "cogs_days": len(parsed["cogs"]), "sales_days": len(parsed["sales"]), "orders_rows": len(parsed["orders"]), "history_snapshots": 1}
    if result != expected:
        raise RuntimeError(f"Post-import counts differ. Expected {expected}, received {result}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
